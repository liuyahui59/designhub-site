#!/usr/bin/env python3
"""Generate a v7 PC B2B UX walkthrough DOCX from report.json."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
except ImportError as exc:  # pragma: no cover - dependency guard
    Document = None  # type: ignore[assignment]
    WD_ALIGN_PARAGRAPH = None  # type: ignore[assignment]
    Inches = None  # type: ignore[assignment]
    Pt = None  # type: ignore[assignment]
    DOCX_IMPORT_ERROR = exc
else:
    DOCX_IMPORT_ERROR = None

MAX_IMAGE_WIDTH_INCHES = 6.2


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="根据 v7 report.json 生成 PC B 端体验走查 DOCX")
    parser.add_argument("--report-json", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--skip-annotation-check", action="store_true", help="仅调试")
    parser.add_argument("--skip-checklist-check", action="store_true", help="仅调试")
    parser.add_argument("--json", action="store_true")
    return parser


def _fail(message: str, detail: str = "", *, reason: str = "tool_error", extra: dict | None = None) -> dict:
    result = {
        "detail": detail[:800],
        "failure_reason": reason,
        "failure_stage": "docx_generation",
        "message": message,
        "ok": False,
        "status": "failed",
    }
    if extra:
        result.update(extra)
    return result


def _text(value: object) -> str:
    if value is None:
        return ""
    return value.strip() if isinstance(value, str) else str(value).strip()


def _string_list(value: object) -> list[str]:
    if not isinstance(value, list):
        return []
    result: list[str] = []
    for item in value:
        text = _text(item)
        if text:
            result.append(text)
    return result


def _object_list(value: object) -> list[dict]:
    return [item for item in value if isinstance(item, dict)] if isinstance(value, list) else []


def _load_json(path: Path) -> dict:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ValueError(f"未找到 report.json：{path}") from exc
    except json.JSONDecodeError as exc:
        raise ValueError(f"report.json 不是合法 JSON：{exc}") from exc
    except OSError as exc:
        raise ValueError(f"读取 report.json 失败：{exc}") from exc
    if not isinstance(data, dict):
        raise ValueError("report.json 根节点必须是对象")
    return data


def _resolve_output_path(raw: str) -> Path:
    output = Path(raw).expanduser().resolve()
    if output.suffix.lower() != ".docx":
        raise ValueError("--output 必须以 .docx 结尾")
    return output


def _required(mapping: dict, key: str, label: str) -> str:
    value = _text(mapping.get(key))
    if not value:
        raise ValueError(f"缺少 {label}")
    return value


def _normalize_review_brief(data: dict) -> dict:
    brief = data.get("review_brief")
    if not isinstance(brief, dict):
        raise ValueError("缺少 review_brief")
    if _text(brief.get("status")).lower() != "confirmed" or brief.get("answered_by_user") is not True:
        raise ValueError("review_brief 必须由用户回答并标记 confirmed")
    context = brief.get("context")
    if not isinstance(context, dict):
        raise ValueError("review_brief.context 必须是对象")
    raw_tasks = brief.get("critical_tasks")
    tasks = _object_list(raw_tasks)
    if not isinstance(raw_tasks, list) or not tasks or len(tasks) != len(raw_tasks):
        raise ValueError("review_brief.critical_tasks 必须是非空任务对象数组")
    normalized_tasks: list[dict] = []
    for index, task in enumerate(tasks):
        task_id = _required(task, "id", f"review_brief.critical_tasks[{index}].id")
        criteria = _string_list(task.get("success_criteria"))
        if not criteria:
            raise ValueError(f"review_brief task {task_id} 缺少成功标准")
        normalized_tasks.append(
            {
                "id": task_id,
                "name": _required(task, "name", f"review_brief task {task_id}.name"),
                "primary_user": _required(task, "primary_user", f"review_brief task {task_id}.primary_user"),
                "start_condition": _required(
                    task, "start_condition", f"review_brief task {task_id}.start_condition"
                ),
                "intended_outcome": _required(
                    task, "intended_outcome", f"review_brief task {task_id}.intended_outcome"
                ),
                "success_criteria": criteria,
            }
        )
    return {
        "object_scope_purpose": _required(brief, "object_scope_purpose", "review_brief.object_scope_purpose"),
        "primary_users": _required(brief, "primary_users", "review_brief.primary_users"),
        "user_goal": _required(brief, "user_goal", "review_brief.user_goal"),
        "critical_tasks": normalized_tasks,
        "context": {
            "frequency": _required(context, "frequency", "review_brief.context.frequency"),
            "display_environment": _required(
                context, "display_environment", "review_brief.context.display_environment"
            ),
            "data_scale": _required(context, "data_scale", "review_brief.context.data_scale"),
            "time_pressure": _required(context, "time_pressure", "review_brief.context.time_pressure"),
            "interruption_collaboration": _required(
                context, "interruption_collaboration", "review_brief.context.interruption_collaboration"
            ),
        },
        "failure_impact": _required(brief, "failure_impact", "review_brief.failure_impact"),
        "fixed_constraints": _required(brief, "fixed_constraints", "review_brief.fixed_constraints"),
    }


def _resolve_image(entry: dict, base_dir: Path) -> dict:
    raw = _required(entry, "path", "images[].path")
    path = Path(raw).expanduser()
    if not path.is_absolute():
        path = (base_dir / path).resolve()
    if not path.is_file():
        raise ValueError(f"图片不存在：{path}")
    return {"path": path, "caption": _text(entry.get("caption"))}


def _normalize_report(data: dict, report_path: Path) -> dict:
    issues: list[dict] = []
    for index, item in enumerate(_object_list(data.get("issues")), start=1):
        issue = dict(item)
        issue["index"] = int(item.get("index", index))
        issue["images"] = [_resolve_image(image, report_path.parent) for image in _object_list(item.get("images"))]
        issues.append(issue)

    inventory = data.get("scope_inventory") if isinstance(data.get("scope_inventory"), dict) else {}
    summary = data.get("summary") if isinstance(data.get("summary"), dict) else {}
    appendix = data.get("appendix") if isinstance(data.get("appendix"), dict) else {}
    review_mode = _text(data.get("review_mode")) or "initial"
    follow_up_review = data.get("follow_up_review") if isinstance(data.get("follow_up_review"), dict) else {}
    cross_page_consistency = (
        data.get("cross_page_consistency") if isinstance(data.get("cross_page_consistency"), dict) else {}
    )
    pc_environment_validation = (
        data.get("pc_environment_validation")
        if isinstance(data.get("pc_environment_validation"), dict)
        else {}
    )
    return {
        "title": _required(data, "title", "title"),
        "project": _text(data.get("project")),
        "source": _required(data, "source", "source"),
        "review_mode": review_mode,
        "report_status": _required(data, "report_status", "report_status"),
        "generated_at": _text(data.get("generated_at")),
        "scope": _text(data.get("scope")),
        "review_brief": _normalize_review_brief(data),
        "activated_packs": _string_list(data.get("activated_packs")),
        "inventory_basis": _object_list(inventory.get("inventory_basis")),
        "inventory_completeness": _text(inventory.get("inventory_completeness")),
        "completeness_reason": _text(inventory.get("completeness_reason")),
        "excluded_surfaces": _object_list(inventory.get("excluded_surfaces")),
        "surfaces": _object_list(inventory.get("surfaces")),
        "tasks": _object_list(inventory.get("tasks")),
        "pc_environment_validation": pc_environment_validation,
        "cross_page_consistency": cross_page_consistency,
        "summary": {
            "overall": _text(summary.get("overall")),
            "key_points": _string_list(summary.get("key_points")),
            "improvement_direction": _text(summary.get("improvement_direction")),
        },
        "issues": issues,
        "validation_backlog": _object_list(data.get("validation_backlog")),
        "positive_patterns": data.get("positive_patterns", []) if isinstance(data.get("positive_patterns", []), list) else [],
        "opportunities": data.get("opportunities", []) if isinstance(data.get("opportunities", []), list) else [],
        "follow_up_review": follow_up_review,
        "appendix_note": _text(appendix.get("note")),
        "appendix_items": _string_list(appendix.get("items")),
        "limitations": _string_list(appendix.get("limitations")),
    }


def _set_base_style(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "PingFang SC"
    normal.font.size = Pt(10.5)


def _add_field(document: Document, label: str, value: object) -> None:
    text = _text(value)
    if not text:
        return
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{label}：").bold = True
    paragraph.add_run(text)


def _add_list(document: Document, heading: str, values: list[str]) -> None:
    if not values:
        return
    document.add_heading(heading, level=2)
    for value in values:
        document.add_paragraph(value, style="List Bullet")


def _add_title(document: Document, report: dict) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Title"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(report["title"])
    run.font.size = Pt(18)
    for line in (
        f"项目：{report['project'] or '未提供'}",
        f"走查类型：{report['review_mode']}",
        f"报告状态：{report['report_status']}",
        f"输入模式：{report['source']}",
        f"走查时间：{report['generated_at'] or '未提供'}",
        f"走查范围：{report['scope'] or '未提供'}",
    ):
        meta = document.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta.add_run(line)


def _add_review_brief(document: Document, report: dict) -> None:
    brief = report["review_brief"]
    context = brief["context"]
    document.add_heading("走查前提", level=1)
    _add_field(document, "对象、范围与目的", brief["object_scope_purpose"])
    _add_field(document, "主要用户", brief["primary_users"])
    _add_field(document, "用户目标", brief["user_goal"])
    _add_field(
        document,
        "PC 使用情境",
        "；".join(
            [
                f"频率：{context['frequency']}",
                f"屏幕与窗口：{context['display_environment']}",
                f"数据量：{context['data_scale']}",
                f"时间压力：{context['time_pressure']}",
                f"中断/协作：{context['interruption_collaboration']}",
            ]
        ),
    )
    document.add_heading("关键任务与成功标准", level=2)
    for task in brief["critical_tasks"]:
        task_heading = document.add_paragraph()
        task_heading.add_run(f"{task['id']}｜{task['name']}").bold = True
        _add_field(document, "主要用户", task["primary_user"])
        _add_field(document, "任务边界", f"{task['start_condition']} → {task['intended_outcome']}")
        _add_list(document, "成功标准", task["success_criteria"])
    _add_field(document, "失败影响", brief["failure_impact"])
    _add_field(document, "业务约束", brief["fixed_constraints"])


def _add_coverage(document: Document, report: dict) -> None:
    document.add_heading("范围与覆盖", level=1)
    _add_field(document, "启用模块", "、".join(report["activated_packs"]))
    _add_field(document, "页面清单完整性", report["inventory_completeness"])
    _add_field(document, "完整性理由", report["completeness_reason"])
    _add_list(
        document,
        "页面清单来源",
        [
            f"{_text(item.get('id'))}｜{_text(item.get('type'))}：{_text(item.get('description'))}"
            for item in report["inventory_basis"]
        ],
    )
    if report["excluded_surfaces"]:
        _add_list(
            document,
            "已知范围外页面",
            [
                f"{_text(item.get('name'))}｜{_text(item.get('reason_code'))}：{_text(item.get('reason'))}"
                for item in report["excluded_surfaces"]
            ],
        )
    else:
        _add_field(document, "已知范围外页面", "无")

    surfaces = report["surfaces"]
    tasks = report["tasks"]
    surface_completed = sum(_text(item.get("visual_status")) == "completed" for item in surfaces)
    task_completed = sum(_text(item.get("flow_status")) == "completed" for item in tasks)
    _add_field(document, "页面视觉覆盖", f"{surface_completed}/{len(surfaces)}")
    _add_field(document, "关键任务覆盖", f"{task_completed}/{len(tasks)}")
    pc_validation = report["pc_environment_validation"]
    environments = _object_list(pc_validation.get("environments"))
    primary = next((item for item in environments if _text(item.get("role")) == "primary"), None)
    if primary:
        zoom = primary.get("browser_zoom_percent")
        zoom_text = f"{zoom}%" if isinstance(zoom, int) and not isinstance(zoom, bool) else "未知"
        environment_detail = "；".join(
            [
                f"{primary.get('viewport_width')}×{primary.get('viewport_height')}",
                _text(primary.get("window_mode")) or "窗口方式未知",
                f"浏览器缩放 {zoom_text}",
                f"覆盖 {len(_string_list(primary.get('surface_ids')))}/{len(surfaces)} 页面",
            ]
        )
        limitation = _text(primary.get("limitations"))
        if limitation:
            environment_detail = f"{environment_detail}；限制：{limitation}"
    else:
        environment_detail = (
            _text(pc_validation.get("reason"))
            or _text(pc_validation.get("target_environment"))
            or "未完成"
        )
    _add_field(
        document,
        "主要 PC 环境",
        f"{_text(pc_validation.get('status'))}｜{environment_detail}",
    )
    cross_page = report["cross_page_consistency"]
    cross_page_status = _text(cross_page.get("status"))
    cross_page_detail = _text(cross_page.get("observation")) or _text(cross_page.get("reason"))
    _add_field(
        document,
        "跨页面一致性",
        f"{cross_page_status}｜{cross_page_detail}" if cross_page_detail else cross_page_status,
    )

    document.add_heading("页面", level=2)
    surface_table = document.add_table(rows=1, cols=3)
    for cell, value in zip(surface_table.rows[0].cells, ("页面", "视觉状态", "证据数")):
        cell.text = value
    for item in surfaces:
        cells = surface_table.add_row().cells
        cells[0].text = _text(item.get("name")) or _text(item.get("id"))
        cells[1].text = _text(item.get("visual_status"))
        cells[2].text = str(len(_string_list(item.get("evidence_ids"))))

    document.add_heading("关键任务", level=2)
    task_table = document.add_table(rows=1, cols=4)
    for cell, value in zip(
        task_table.rows[0].cells,
        ("任务", "流程状态", "状态覆盖摘要", "证据数"),
    ):
        cell.text = value
    for item in tasks:
        cells = task_table.add_row().cells
        state_inventory = item.get("state_inventory") if isinstance(item.get("state_inventory"), dict) else {}
        state_results = _object_list(state_inventory.get("results"))
        status_labels = {
            "verified_pass": "通过",
            "issue": "问题",
            "not_tested": "未验证",
            "blocked": "受阻",
        }
        grouped_states: list[str] = []
        for status, label in status_labels.items():
            names = [
                _text(result.get("state_name")) or _text(result.get("state_id"))
                for result in state_results
                if _text(result.get("status")) == status
            ]
            if names:
                grouped_states.append(f"{label}：{'、'.join(names)}")
        state_summary = (
            f"{_text(state_inventory.get('inventory_completeness'))}｜"
            f"{sum(_text(result.get('status')) in {'verified_pass', 'issue'} for result in state_results)}"
            f"/{len(state_results)} 已形成结论"
        )
        if grouped_states:
            state_summary += "\n" + "；".join(grouped_states)
        cells[0].text = _text(item.get("name")) or _text(item.get("id"))
        cells[1].text = _text(item.get("flow_status"))
        cells[2].text = state_summary
        cells[3].text = str(len(_string_list(item.get("evidence_ids"))))


def _add_summary(document: Document, report: dict) -> None:
    summary = report["summary"]
    document.add_heading("整体评价", level=1)
    document.add_paragraph(summary["overall"] or "未提供整体评价。")
    _add_list(document, "核心结论", summary["key_points"])
    _add_field(document, "改进方向", summary["improvement_direction"])


def _acceptance_result_line(item: dict) -> str:
    result_labels = {"passed": "通过", "failed": "未通过", "not_verified": "未验证"}
    result = _text(item.get("result"))
    details = _text(item.get("observation")) or _text(item.get("reason"))
    evidence = "、".join(_string_list(item.get("evidence_ids")))
    validation = _text(item.get("validation_method"))
    parts = [details] if details else []
    if evidence:
        parts.append(f"证据 {evidence}")
    if validation:
        parts.append(f"验证方式 {validation}")
    suffix = f"：{'；'.join(parts)}" if parts else ""
    return f"[{result_labels.get(result, result)}] {_text(item.get('criterion'))}{suffix}"


def _add_follow_up_review(document: Document, report: dict) -> None:
    if report["review_mode"] != "follow_up":
        return
    follow_up = report["follow_up_review"]
    baseline = follow_up.get("baseline_report") if isinstance(follow_up.get("baseline_report"), dict) else {}
    document.add_heading("设计复查结论", level=1)
    _add_field(document, "原报告", _text(baseline.get("title")))
    _add_field(document, "原报告时间", baseline.get("generated_at"))
    _add_field(document, "原报告来源", baseline.get("source_reference"))
    _add_field(document, "前提变化", follow_up.get("brief_change_status"))
    _add_list(document, "前提变化项", _string_list(follow_up.get("brief_changes")))
    _add_field(document, "复查范围", follow_up.get("scope_type"))

    for item in _object_list(follow_up.get("issue_results")):
        title = _text(item.get("title")) or _text(item.get("original_issue_id"))
        document.add_heading(title, level=2)
        _add_field(document, "原问题 ID", item.get("original_issue_id"))
        _add_field(document, "复查状态", item.get("status"))
        _add_field(document, "关联任务", "、".join(_string_list(item.get("task_ids"))))
        _add_field(document, "关联页面", "、".join(_string_list(item.get("surface_ids"))))
        _add_list(
            document,
            "验收标准结果",
            [_acceptance_result_line(entry) for entry in _object_list(item.get("acceptance_results"))],
        )
        _add_field(document, "当前问题 ID", item.get("current_issue_id"))
        _add_field(document, "结论", item.get("conclusion"))

    regression = follow_up.get("regression_check") if isinstance(follow_up.get("regression_check"), dict) else {}
    document.add_heading("回归检查", level=2)
    _add_field(document, "结果", regression.get("result"))
    _add_field(document, "覆盖任务", "、".join(_string_list(regression.get("task_ids"))))
    _add_field(document, "覆盖页面", "、".join(_string_list(regression.get("surface_ids"))))
    _add_field(document, "观察", regression.get("observation"))
    _add_field(document, "证据", "、".join(_string_list(regression.get("evidence_ids"))))


def _issue_heading(issue: dict) -> str:
    return f"{issue['index']}. [{_text(issue.get('severity')).upper()}] {_text(issue.get('title'))}"


def _occurrence_line(item: dict) -> str:
    scopes: list[str] = []
    if _string_list(item.get("task_ids")):
        scopes.append("任务 " + "、".join(_string_list(item.get("task_ids"))))
    if _string_list(item.get("surface_ids")):
        scopes.append("页面 " + "、".join(_string_list(item.get("surface_ids"))))
    if _text(item.get("state")):
        scopes.append("状态 " + _text(item.get("state")))
    if _string_list(item.get("evidence_ids")):
        scopes.append("证据 " + "、".join(_string_list(item.get("evidence_ids"))))
    prefix = f"{_text(item.get('id'))}｜" if _text(item.get("id")) else ""
    return f"{prefix}{'；'.join(scopes)}：{_text(item.get('observation'))}"


def _add_issue(document: Document, issue: dict) -> int:
    document.add_heading(_issue_heading(issue), level=2)
    if _text(issue.get("issue_origin")):
        _add_field(document, "问题来源", issue.get("issue_origin"))
    if _text(issue.get("original_issue_id")):
        _add_field(document, "原问题 ID", issue.get("original_issue_id"))
    _add_field(document, "受影响用户", "、".join(_string_list(issue.get("affected_users"))))
    _add_field(document, "关键任务", "、".join(_string_list(issue.get("task_ids"))))
    _add_field(document, "页面", "、".join(_string_list(issue.get("surface_ids"))))
    _add_field(document, "状态", issue.get("state"))
    _add_field(document, "当前观察", issue.get("observation"))
    _add_field(document, "预期体验", issue.get("expected_experience"))
    _add_field(document, "用户影响", issue.get("user_impact"))
    _add_field(document, "业务影响", issue.get("business_impact"))
    _add_field(document, "影响依据", issue.get("impact_basis"))
    _add_field(document, "影响依据理由", issue.get("impact_basis_rationale"))
    _add_field(document, "影响证据", "、".join(_string_list(issue.get("impact_evidence_ids"))))
    _add_list(document, "复现步骤", _string_list(issue.get("reproduction")))
    _add_field(document, "严重度与置信度", f"{_text(issue.get('severity')).upper()} / {_text(issue.get('confidence'))}")
    _add_field(document, "定级理由", issue.get("severity_reason"))
    _add_list(document, "发生实例", [_occurrence_line(item) for item in _object_list(issue.get("occurrences"))])
    _add_field(document, "改进目标", issue.get("design_objective"))
    _add_list(document, "建议方向", _string_list(issue.get("recommendation")))
    _add_field(document, "推荐理由", issue.get("recommendation_rationale"))
    _add_list(document, "约束与权衡", _string_list(issue.get("constraints_tradeoffs")))
    _add_list(document, "备选方向", _string_list(issue.get("alternatives")))
    _add_list(document, "验收标准", _string_list(issue.get("acceptance_criteria")))

    image_count = 0
    for image in issue["images"]:
        document.add_picture(str(image["path"]), width=Inches(MAX_IMAGE_WIDTH_INCHES))
        image_count += 1
        if image["caption"]:
            caption = document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.add_run(image["caption"]).italic = True
    return image_count


def _add_issues(document: Document, report: dict) -> int:
    issues = report["issues"]
    document.add_heading("确定问题", level=1)
    if not issues:
        document.add_paragraph("在本次用户、任务、场景与证据范围内，未发现达到正式问题标准的确定缺陷。")
        return 0

    headings = {
        "task": "关键任务体验问题",
        "visual": "页面视觉质量问题",
        "visual_system": "跨页面系统性问题",
    }
    total_images = 0
    for track in ("task", "visual", "visual_system"):
        track_issues = [issue for issue in issues if _text(issue.get("evaluation_track")) == track]
        if not track_issues:
            continue
        document.add_heading(headings[track], level=1)
        for issue in track_issues:
            total_images += _add_issue(document, issue)
    return total_images


def _add_backlog(document: Document, report: dict) -> None:
    document.add_heading("待验证风险", level=1)
    risks = report["validation_backlog"]
    if not risks:
        document.add_paragraph("无。")
        return
    table = document.add_table(rows=1, cols=4)
    for cell, value in zip(table.rows[0].cells, ("风险", "潜在影响", "验证优先级", "验证方式")):
        cell.text = value
    for risk in risks:
        cells = table.add_row().cells
        cells[0].text = _text(risk.get("title"))
        cells[1].text = _text(risk.get("potential_impact"))
        cells[2].text = _text(risk.get("validation_priority"))
        cells[3].text = _text(risk.get("validation_method"))


def _add_positive_patterns(document: Document, items: list[object]) -> None:
    document.add_heading("正向设计模式", level=1)
    patterns = _object_list(items)
    if not patterns:
        document.add_paragraph("无。")
        return
    for item in patterns:
        document.add_heading(_text(item.get("title")) or _text(item.get("id")), level=2)
        _add_field(document, "具体观察", item.get("observation"))
        _add_field(document, "用户价值", item.get("user_value"))
        _add_field(document, "保留与复用指导", item.get("preservation_guidance"))
        _add_field(document, "关联任务", "、".join(_string_list(item.get("task_ids"))))
        _add_field(document, "关联页面", "、".join(_string_list(item.get("surface_ids"))))
        _add_field(document, "检查项", "、".join(_string_list(item.get("checklist_ids"))))
        _add_field(document, "证据", "、".join(_string_list(item.get("evidence_ids"))))


def _add_opportunities(document: Document, items: list[object]) -> None:
    document.add_heading("可选增强机会", level=1)
    opportunities = _object_list(items)
    if not opportunities:
        document.add_paragraph("无。")
        return
    for item in opportunities:
        document.add_heading(_text(item.get("title")) or _text(item.get("id")), level=2)
        _add_field(document, "当前状态", item.get("current_state"))
        _add_field(document, "为什么不构成问题", item.get("why_not_issue"))
        _add_field(document, "预期价值", item.get("expected_value"))
        _add_field(document, "适用条件", item.get("applicability"))
        _add_field(document, "设计方向", item.get("design_direction"))
        _add_field(document, "关联任务", "、".join(_string_list(item.get("task_ids"))))
        _add_field(document, "关联页面", "、".join(_string_list(item.get("surface_ids"))))
        _add_field(document, "检查项", "、".join(_string_list(item.get("checklist_ids"))))
        _add_field(document, "证据", "、".join(_string_list(item.get("evidence_ids"))))


def _add_appendix(document: Document, report: dict, total_images: int) -> None:
    document.add_heading("未覆盖范围与证据限制", level=1)
    if report["limitations"]:
        for item in report["limitations"]:
            document.add_paragraph(item, style="List Bullet")
    else:
        document.add_paragraph("无已知限制。")
    document.add_heading("附录", level=1)
    _add_field(document, "说明", report["appendix_note"])
    for item in report["appendix_items"]:
        document.add_paragraph(item, style="List Bullet")
    if total_images == 0:
        document.add_paragraph("本次报告无问题标注图。")


def _generate_docx(report: dict, output_path: Path) -> dict:
    if Document is None or WD_ALIGN_PARAGRAPH is None or Inches is None or Pt is None:
        return _fail("DOCX generation failed: python-docx not available", detail=f"缺少依赖：{DOCX_IMPORT_ERROR}")
    document = Document()
    _set_base_style(document)
    _add_title(document, report)
    _add_review_brief(document, report)
    _add_coverage(document, report)
    _add_summary(document, report)
    _add_follow_up_review(document, report)
    total_images = _add_issues(document, report)
    _add_backlog(document, report)
    _add_positive_patterns(document, report["positive_patterns"])
    _add_opportunities(document, report["opportunities"])
    _add_appendix(document, report, total_images)
    try:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output_path))
    except Exception as exc:  # pragma: no cover - library specific
        return _fail("DOCX generation failed", detail=f"写入 DOCX 失败：{exc}")
    return {
        "file_name": output_path.name,
        "image_count": total_images,
        "issue_count": len(report["issues"]),
        "message": "DOCX generated",
        "ok": True,
        "output_path": str(output_path),
        "report_status": report["report_status"],
        "review_mode": report["review_mode"],
        "source": report["source"],
        "status": "generated",
        "validation_backlog_count": len(report["validation_backlog"]),
    }


def build_result(args: argparse.Namespace) -> dict:
    try:
        output_path = _resolve_output_path(args.output)
    except ValueError as exc:
        return _fail("DOCX generation failed: invalid output", detail=str(exc))
    if output_path.exists() and not args.overwrite:
        return _fail("DOCX generation failed: output exists", detail=f"输出文件已存在：{output_path}")

    report_path = Path(args.report_json).expanduser().resolve()
    try:
        raw = _load_json(report_path)
    except ValueError as exc:
        return _fail("DOCX generation failed: invalid report json", detail=str(exc))

    if not args.skip_checklist_check:
        try:
            from validate_checklist_coverage import validate_report
        except ImportError as exc:
            return _fail("DOCX generation failed: report validator missing", detail=str(exc))
        gate = validate_report(raw)
        if not gate.get("ok"):
            errors = gate.get("errors", [])
            return _fail(
                "DOCX generation failed: report coverage gate not passed",
                detail="；".join(errors[:8]),
                reason="report_structure_invalid",
                extra={"report_errors": errors},
            )

    if not args.skip_annotation_check:
        try:
            from validate_report_annotations import validate_report as validate_annotations
        except ImportError as exc:
            return _fail("DOCX generation failed: annotation validator missing", detail=str(exc))
        gate = validate_annotations(raw, base_dir=report_path.parent)
        if not gate.get("ok"):
            errors = gate.get("errors", [])
            return _fail(
                "DOCX generation failed: annotation gate not passed",
                detail="；".join(errors[:8]),
                reason="annotations_missing",
                extra={"annotation_errors": errors},
            )

    try:
        report = _normalize_report(raw, report_path)
    except ValueError as exc:
        return _fail("DOCX generation failed: invalid report json", detail=str(exc))
    return _generate_docx(report, output_path)


def main() -> int:
    args = _build_parser().parse_args()
    result = build_result(args)
    if args.json or not result.get("ok"):
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print(result["output_path"])
    return 0 if result.get("ok") else 1


if __name__ == "__main__":
    raise SystemExit(main())
